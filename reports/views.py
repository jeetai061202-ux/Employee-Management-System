from datetime import datetime

from django.contrib.auth.decorators import login_required
from django.db.models import Count
from django.http import HttpResponse
from django.shortcuts import render

from openpyxl import Workbook

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch

from accounts.decorators import role_required

from employees.models import Employee
from attendance.models import Attendance
from payroll.models import Payroll


# ==========================================================
# Reports Dashboard
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def report_dashboard(request):

    return render(
        request,
        "reports/report_dashboard.html"
    )


# ==========================================================
# Employee Report
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def employee_report(request):

    employees = Employee.objects.all().order_by("-id")[:10]

    department_data = (
        Employee.objects
        .values("department")
        .annotate(total=Count("id"))
        .order_by("department")
    )

    context = {

        "employees": employees,

        "total_employees": Employee.objects.count(),

        "active_employees": Employee.objects.filter(
            is_active=True
        ).count(),

        "inactive_employees": Employee.objects.filter(
            is_active=False
        ).count(),

        "department_labels": [
            item["department"]
            for item in department_data
        ],

        "department_counts": [
            item["total"]
            for item in department_data
        ],

    }

    return render(
        request,
        "reports/employee_report.html",
        context
    )


# ==========================================================
# Payroll Report
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def payroll_report(request):

    payrolls = Payroll.objects.select_related(
        "employee"
    ).order_by("-year", "-month")

    context = {

        "payrolls": payrolls[:20],

        "total_payrolls": Payroll.objects.count(),

        "total_net_salary": sum(
            payroll.net_salary
            for payroll in payrolls
        ),

        "total_basic_salary": sum(
            payroll.basic_salary
            for payroll in payrolls
        ),

        "total_bonus": sum(
            payroll.bonus
            for payroll in payrolls
        ),

        "chart_labels": [
            "Basic Salary",
            "Bonus",
            "Net Salary",
        ],

        "chart_values": [

            sum(
                payroll.basic_salary
                for payroll in payrolls
            ),

            sum(
                payroll.bonus
                for payroll in payrolls
            ),

            sum(
                payroll.net_salary
                for payroll in payrolls
            ),

        ],

    }

    return render(
        request,
        "reports/payroll_report.html",
        context
    )


# ==========================================================
# Attendance Report
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def attendance_report(request):

    attendance = Attendance.objects.select_related(
        "employee"
    ).order_by("-date")[:20]

    context = {

        "attendance": attendance,

        "total_records": Attendance.objects.count(),

        "present_count": Attendance.objects.filter(
            status="Present"
        ).count(),

        "leave_count": Attendance.objects.filter(
            status="Leave"
        ).count(),

        "halfday_count": Attendance.objects.filter(
            status="Half Day"
        ).count(),

        "chart_labels": [
            "Present",
            "Leave",
            "Half Day",
        ],

        "chart_values": [

            Attendance.objects.filter(
                status="Present"
            ).count(),

            Attendance.objects.filter(
                status="Leave"
            ).count(),

            Attendance.objects.filter(
                status="Half Day"
            ).count(),

        ],

    }

    return render(
        request,
        "reports/attendance_report.html",
        context
    )


# ==========================================================
# Export Employee Report - Excel
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def employee_report_excel(request):

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Employee Report"

    worksheet.append([
        "Employee ID",
        "First Name",
        "Last Name",
        "Department",
        "Designation",
        "Status",
        "Email",
        "Phone",
    ])

    employees = Employee.objects.all()

    for employee in employees:

        worksheet.append([

            employee.employee_id,

            employee.first_name,

            employee.last_name,

            employee.department,

            employee.designation,

            "Active"
            if employee.is_active
            else "Inactive",

            employee.email,

            employee.phone,

        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="Employee_Report.xlsx"'
    )

    workbook.save(response)

    return response


# ==========================================================
# Export Employee Report - PDF
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def employee_report_pdf(request):

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="Employee_Report.pdf"'
    )

    doc = SimpleDocTemplate(response)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Employee Management System</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            f"Generated on "
            f"{datetime.now().strftime('%d-%m-%Y %H:%M')}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            0.3 * inch
        )
    )

    data = [[
        "Employee ID",
        "Employee Name",
        "Department",
        "Designation",
        "Status",
    ]]

    employees = Employee.objects.all()

    for employee in employees:

        data.append([

            employee.employee_id,

            (
                f"{employee.first_name} "
                f"{employee.last_name}"
            ),

            employee.department,

            employee.designation,

            "Active"
            if employee.is_active
            else "Inactive",

        ])

    table = Table(data)

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.darkblue
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                1,
                colors.black
            ),

            (
                "BACKGROUND",
                (0, 1),
                (-1, -1),
                colors.beige
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

        ])
    )

    elements.append(table)

    doc.build(elements)

    return response


# ==========================================================
# Export Attendance Report - Excel
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def attendance_report_excel(request):

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Attendance Report"

    worksheet.append([
        "Employee ID",
        "Employee Name",
        "Date",
        "Status",
        "Check In",
        "Check Out",
        "Working Hours",
    ])

    attendance = Attendance.objects.select_related(
        "employee"
    ).all()

    for record in attendance:

        worksheet.append([

            record.employee.employee_id,

            (
                f"{record.employee.first_name} "
                f"{record.employee.last_name}"
            ),

            record.date,

            record.status,

            str(record.check_in),

            str(record.check_out),

            record.working_hours,

        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="Attendance_Report.xlsx"'
    )

    workbook.save(response)

    return response


# ==========================================================
# Export Attendance Report - PDF
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def attendance_report_pdf(request):

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="Attendance_Report.pdf"'
    )

    doc = SimpleDocTemplate(response)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Attendance Report</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            f"Generated on "
            f"{datetime.now().strftime('%d-%m-%Y %H:%M')}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            0.3 * inch
        )
    )

    data = [[
        "Employee",
        "Date",
        "Status",
        "Check In",
        "Check Out",
    ]]

    attendance = Attendance.objects.select_related(
        "employee"
    ).all()

    for record in attendance:

        data.append([

            (
                f"{record.employee.first_name} "
                f"{record.employee.last_name}"
            ),

            str(record.date),

            record.status,

            str(record.check_in),

            str(record.check_out),

        ])

    table = Table(data)

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.darkblue
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                1,
                colors.black
            ),

            (
                "BACKGROUND",
                (0, 1),
                (-1, -1),
                colors.beige
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

        ])
    )

    elements.append(table)

    doc.build(elements)

    return response


# ==========================================================
# Export Payroll Report - Excel
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def payroll_report_excel(request):

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Payroll Report"

    worksheet.append([
        "Employee ID",
        "Employee Name",
        "Month",
        "Year",
        "Basic Salary",
        "Bonus",
        "Tax",
        "PF",
        "Net Salary",
    ])

    payrolls = Payroll.objects.select_related(
        "employee"
    ).all()

    for payroll in payrolls:

        worksheet.append([

            payroll.employee.employee_id,

            (
                f"{payroll.employee.first_name} "
                f"{payroll.employee.last_name}"
            ),

            payroll.month,

            payroll.year,

            payroll.basic_salary,

            payroll.bonus,

            payroll.tax,

            payroll.pf,

            payroll.net_salary,

        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="Payroll_Report.xlsx"'
    )

    workbook.save(response)

    return response


# ==========================================================
# Export Payroll Report - PDF
# ==========================================================

@login_required
@role_required("ADMIN", "HR")
def payroll_report_pdf(request):

    response = HttpResponse(
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        'attachment; filename="Payroll_Report.pdf"'
    )

    doc = SimpleDocTemplate(response)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b>Payroll Report</b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            f"Generated on "
            f"{datetime.now().strftime('%d-%m-%Y %H:%M')}",
            styles["Normal"]
        )
    )

    elements.append(
        Spacer(
            1,
            0.3 * inch
        )
    )

    data = [[
        "Employee",
        "Month",
        "Year",
        "Basic",
        "Bonus",
        "Net Salary",
    ]]

    payrolls = Payroll.objects.select_related(
        "employee"
    ).all()

    for payroll in payrolls:

        data.append([

            (
                f"{payroll.employee.first_name} "
                f"{payroll.employee.last_name}"
            ),

            payroll.month,

            payroll.year,

            f"₹ {payroll.basic_salary}",

            f"₹ {payroll.bonus}",

            f"₹ {payroll.net_salary}",

        ])

    table = Table(data)

    table.setStyle(
        TableStyle([

            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.darkblue
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                1,
                colors.black
            ),

            (
                "BACKGROUND",
                (0, 1),
                (-1, -1),
                colors.beige
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),

        ])
    )

    elements.append(table)

    doc.build(elements)

    return response